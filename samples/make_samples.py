"""Create messy sample files for testing: title rows, total rows, Excel dates, numeric text.

Run:  python samples/make_samples.py
"""
import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

REGIONS = ["Lahore", "Karachi", "Islamabad", "Multan"]
PRODUCTS = {"Tea": 450, "Sugar": 160, "Rice": 320, "Oil": 610, "Flour": 120}
SALESPEOPLE = ["Ali", "Sara", "Usman", "Hina", "Bilal"]


def sales_rows(seed: int, start: date, days: int, n: int):
    rnd = random.Random(seed)
    rows = []
    for i in range(n):
        product = rnd.choice(list(PRODUCTS))
        qty = rnd.randint(1, 40)
        price = PRODUCTS[product]
        rows.append({
            "order_id": 1000 + seed * 1000 + i,
            "date": start + timedelta(days=rnd.randrange(days)),
            "region": rnd.choice(REGIONS),
            "product": product,
            "salesperson": rnd.choice(SALESPEOPLE),
            "qty": qty,
            "unit_price": price,
            "amount": qty * price,
        })
    # one obvious outlier
    rows[n // 2]["qty"] = 400
    rows[n // 2]["amount"] = 400 * rows[n // 2]["unit_price"]
    return rows


def make_sales_report(folder: Path) -> Path:
    """Title rows, an empty column, real Excel dates, some amounts typed as text, a Total row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Data"
    ws.append(["ABC Traders — Sales Report"])
    ws.append(["Period: Jan–Mar 2026"])
    ws.append([])
    ws.append(["Order ID", "Order Date", "Region", None, "Product", "Sales Person", "Qty", "Unit Price", "Amount (Rs)"])
    rows = sales_rows(1, date(2026, 1, 1), 90, 120)
    for i, r in enumerate(rows):
        amount = r["amount"]
        if i % 10 == 3:
            amount = f"{amount:,}"            # "1,200"
        elif i % 10 == 7:
            amount = f"Rs {amount}"          # "Rs 500"
        ws.append([r["order_id"], r["date"], r["region"], None, r["product"], r["salesperson"],
                   r["qty"], r["unit_price"], amount])
        ws.cell(row=ws.max_row, column=2).number_format = "DD-MMM-YYYY"
    ws.append([])
    ws.append(["Total", None, None, None, None, None, sum(r["qty"] for r in rows), None,
               sum(r["amount"] for r in rows)])

    t = wb.create_sheet("Targets")
    t.append(["Region", "Monthly Target"])
    for reg, tgt in zip(REGIONS, [900000, 1100000, 600000, 400000]):
        t.append([reg, tgt])
    t.append(["Grand Total", 3000000])
    path = folder / "sales_report.xlsx"
    wb.save(path)
    return path


def make_monthly_sales(folder: Path) -> Path:
    """Three monthly sheets with the same columns; dates stored as Excel serial numbers; subtotal rows."""
    wb = Workbook()
    wb.remove(wb.active)
    for m, name in enumerate(["Jan", "Feb", "Mar"], start=1):
        ws = wb.create_sheet(name)
        ws.append([f"Branch sales — {name} 2026"])
        ws.append(["Date", "Branch", "Units", "Revenue"])
        rnd = random.Random(m)
        for d in range(1, 29):
            day = date(2026, m, d)
            serial = (day - date(1899, 12, 30)).days       # Excel serial number, no date format
            for branch in ["North", "South"]:
                units = rnd.randint(10, 90)
                ws.append([serial, branch, units, units * 250])
        ws.append(["Subtotal", None, "=SUM(C3:C58)", "=SUM(D3:D58)"])
    path = folder / "monthly_sales.xlsx"
    wb.save(path)
    return path


def make_expenses_csv(folder: Path) -> Path:
    """dd/mm/yyyy date text, amounts with commas, '-' placeholders and a Grand Total row."""
    path = folder / "expenses.csv"
    rnd = random.Random(7)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Expense Report FY2026"])
        w.writerow([])
        w.writerow(["Date", "Category", "Amount", "Paid By"])
        total = 0
        for i in range(40):
            day = date(2026, 1, 1) + timedelta(days=i * 2)
            amount = rnd.randint(500, 25000)
            total += amount
            w.writerow([day.strftime("%d/%m/%Y"), rnd.choice(["Rent", "Fuel", "Salaries", "Utilities"]),
                        f"{amount:,}", "-" if i % 9 == 0 else rnd.choice(["Cash", "Bank"])])
        w.writerow(["Grand Total", "", f"{total:,}", ""])
    return path


STORES = [("Lahore", "Packages Mall"), ("Lahore", "Emporium"), ("Karachi", "Dolmen Clifton"),
          ("Karachi", "Lucky One"), ("Islamabad", "Centaurus"), ("Multan", "Mall of Multan")]
CATEGORIES = ["Unstitched", "Ready-to-wear", "Home textile", "Accessories"]


def store_sales_rows(seed: int = 11):
    rnd = random.Random(seed)
    rows = []
    for month in ["2026-07", "2026-08", "2026-09"]:
        for city, store in STORES:
            for cat in CATEGORIES:
                units = rnd.randint(120, 900)
                rows.append((month, city, store, cat, units, units * rnd.choice([1800, 2400, 3200, 4500])))
    return rows


def make_store_sales_pdf(folder: Path) -> Path:
    """A PDF report with a bordered table that runs over two pages (header repeated) and a total row."""
    from fpdf import FPDF
    rows = store_sales_rows()
    pdf = FPDF()
    pdf.set_auto_page_break(False)
    widths, header = [24, 26, 44, 34, 22, 36], ["Month", "City", "Store", "Category", "Units", "Net Sales (Rs)"]

    def page_header(first: bool):
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14 if first else 10)
        pdf.cell(0, 9, "Retail Store Sales Report - Q3 2026" if first else "Retail Store Sales Report (continued)",
                 new_x="LMARGIN", new_y="NEXT")
        if first:
            pdf.set_font("Helvetica", "", 9)
            pdf.multi_cell(0, 5, "Quarterly sales by store and category. Figures in PKR, net of returns.")
            pdf.ln(2)
        pdf.set_font("Helvetica", "B", 8)
        for w, h in zip(widths, header):
            pdf.cell(w, 7, h, border=1)
        pdf.ln()
        pdf.set_font("Helvetica", "", 8)

    page_header(True)
    for i, r in enumerate(rows):
        if pdf.get_y() > 275:
            page_header(False)
        for w, v in zip(widths, [r[0], r[1], r[2], r[3], f"{r[4]:,}", f"{r[5]:,}"]):
            pdf.cell(w, 6, str(v), border=1)
        pdf.ln()
    pdf.set_font("Helvetica", "B", 8)
    for w, v in zip(widths, ["Total", "", "", "", f"{sum(r[4] for r in rows):,}", f"{sum(r[5] for r in rows):,}"]):
        pdf.cell(w, 6, v, border=1)
    path = folder / "store_sales_q3.pdf"
    pdf.output(str(path))
    return path


def make_text_report_pdf(folder: Path) -> Path:
    """A narrative PDF with figures but no table lines (needs AI extraction)."""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Monthly Management Summary - September 2026", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, "Online sales grew strongly this month while two outlets were under renovation.\n\n"
                         "Channel results for September: Retail stores Rs 48,500,000; E-commerce Rs 12,300,000; "
                         "Wholesale Rs 9,750,000; Export Rs 6,100,000.\n\n"
                         "Returns rate by channel: Retail 3.1%, E-commerce 7.4%, Wholesale 1.2%, Export 0.8%.")
    path = folder / "management_summary_sep.pdf"
    pdf.output(str(path))
    return path


def build_all(folder: Path) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    return [make_sales_report(folder), make_monthly_sales(folder), make_expenses_csv(folder)]


def build_pdfs(folder: Path) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    return [make_store_sales_pdf(folder), make_text_report_pdf(folder)]


if __name__ == "__main__":
    for p in build_all(Path(__file__).parent) + build_pdfs(Path(__file__).parent):
        print("created", p)
